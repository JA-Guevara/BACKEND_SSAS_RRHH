import importlib
import json
import os
from contextlib import asynccontextmanager
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import httpx
import pytest
from docx import Document
from pydantic import ValidationError
from reportlab.pdfgen.canvas import Canvas
from sqlalchemy.ext.asyncio import AsyncEngine

from ssas.analisis_cv.domain.analysis import (
    AnalisisCvError,
    ResultadoIA,
    calcular_afinidad,
    validar_evidencia,
)
from ssas.analisis_cv.infrastructure.extraction.cv_extractor import LocalCvExtractor, extract_file
from ssas.analisis_cv.infrastructure.providers.analysis_lock import analysis_lock
from ssas.analisis_cv.infrastructure.providers.gemini_provider import GeminiAnalysisProvider
from ssas.analisis_cv.infrastructure.providers.privacy import redact_personal_data
from ssas.config.settings import Settings
from ssas.infrastructure.database.base import import_all_models

import_all_models()


def config(**kwargs):
    return Settings(
        _env_file=None,
        app_env="development",
        app_secret_key="test",
        gemini_api_key="test-key",
        **kwargs,
    )


def result(**kwargs):
    return ResultadoIA.model_validate(
        {
            "habilidades": [
                {"habilidad_id": "python", "nivel": "INTERMEDIO", "evidencia": "Desarrollo Python"}
            ],
            "anios_experiencia": 2,
            "evidencia_experiencia": "2 anos de experiencia",
            "resumen": "Perfil tecnico",
            "justificacion": "Evidencia laboral",
            **kwargs,
        }
    )


def test_weighted_formula_and_caps():
    requirements = [{"habilidad_id": "python", "peso": 3}, {"habilidad_id": "sql", "peso": 1}]
    assert calcular_afinidad(requirements, result(), 4) == Decimal("70.00")
    assert calcular_afinidad(requirements[:1], result(), 1) == Decimal("100.00")
    assert calcular_afinidad([], result(), 0) == Decimal("100.00")
    assert calcular_afinidad(requirements, result(habilidades=[], anios_experiencia=0), 2) == 0


@pytest.mark.parametrize("years", [-1, 81, float("inf"), float("nan")])
def test_invalid_experience(years):
    with pytest.raises(ValidationError):
        result(anios_experiencia=years)


def test_catalog_and_literal_evidence():
    text = "Desarrollo Python. 2 anos de experiencia."
    validar_evidencia(result(), {"python"}, text)
    for catalog, document, value in [
        ({"sql"}, text, result()),
        ({"python"}, "Sin evidencia", result()),
        ({"python"}, text, result(habilidades=result().model_dump()["habilidades"] * 2)),
    ]:
        with pytest.raises(AnalisisCvError) as exc:
            validar_evidencia(value, catalog, document)
        assert exc.value.status_code == 502


def test_privacy():
    text = "Juan Perez\nCI: 12345\njuan@example.org\n+591 70012345\nDesarrollo Python"
    cleaned = redact_personal_data(text, ["Juan", "Perez", "12345"])
    assert all(value not in cleaned for value in ["Juan", "Perez", "12345", "example.org"])
    assert "Desarrollo Python" in cleaned


def test_docx_extract_tables_and_limits(tmp_path):
    document = Document()
    document.add_paragraph("Desarrollo Python")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "SQL"
    path = tmp_path / "cv.docx"
    document.save(path)
    text = extract_file(path, 1000000, 1000)
    assert "Desarrollo Python" in text and "SQL" in text
    for bytes_limit, chars_limit in [(1, 1000), (1000000, 3)]:
        with pytest.raises(AnalisisCvError):
            extract_file(path, bytes_limit, chars_limit)


def test_pdf_extract_and_scanned(tmp_path):
    path = tmp_path / "cv.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(40, 700, "Desarrollo Python")
    canvas.save()
    assert "Desarrollo Python" in extract_file(path, 100000, 1000)
    empty = tmp_path / "empty.pdf"
    canvas = Canvas(str(empty))
    canvas.showPage()
    canvas.save()
    with pytest.raises(AnalisisCvError, match="OCR"):
        extract_file(empty, 100000, 1000)


@pytest.mark.asyncio
async def test_extractor_worker_and_storage_boundary(tmp_path):
    document = Document()
    document.add_paragraph("Python")
    document.save(tmp_path / "cv.docx")
    extractor = LocalCvExtractor(config(cv_storage_directory=str(tmp_path)))
    assert await extractor.extract("uploads/cv/cv.docx") == "Python"
    with pytest.raises(AnalisisCvError) as exc:
        await extractor.extract("../../.env")
    assert exc.value.status_code == 404
    (tmp_path / "bad.docx").write_bytes(b"not a document")
    with pytest.raises(AnalisisCvError):
        await extractor.extract("bad.docx")


@pytest.mark.asyncio
async def test_extractor_worker_receives_source_path(tmp_path, monkeypatch):
    module = importlib.import_module("ssas.analisis_cv.infrastructure.extraction.cv_extractor")
    (tmp_path / "cv.pdf").write_bytes(b"sample")
    monkeypatch.setenv("PYTHONPATH", "existing-path")

    def run_worker(args, **kwargs):
        assert args[2] == "ssas.analisis_cv.infrastructure.extraction.worker"
        source_root = str(Path(module.__file__).resolve().parents[4])
        assert kwargs["env"]["PYTHONPATH"].split(os.pathsep) == [source_root, "existing-path"]
        return SimpleNamespace(returncode=0, stdout=b'{"text":"CV de prueba"}')

    monkeypatch.setattr(module.subprocess, "run", run_worker)
    extractor = LocalCvExtractor(config(cv_storage_directory=str(tmp_path)))
    assert await extractor.extract("cv.pdf") == "CV de prueba"


@pytest.mark.asyncio
async def test_provider_structured_request():
    requests = []

    async def handler(request):
        requests.append(json.loads(request.content))
        assert request.headers["x-goog-api-key"] == "test-key"
        assert request.url.path.endswith("/models/gemini-test-model:generateContent")
        return httpx.Response(
            200,
            json={
                "candidates": [
                    {
                        "finishReason": "STOP",
                        "content": {"parts": [{"text": result().model_dump_json()}]},
                    }
                ],
            },
        )

    provider = GeminiAnalysisProvider(
        config(gemini_cv_model="gemini-test-model"), httpx.MockTransport(handler)
    )
    assert await provider.analyze("CV", {}, []) == result()
    assert requests[0]["generationConfig"]["responseMimeType"] == "application/json"
    assert "responseFormat" not in requests[0]["generationConfig"]
    assert json.loads(requests[0]["contents"][0]["parts"][0]["text"])["cv"] == "CV"


@pytest.mark.asyncio
async def test_provider_reports_persistent_bad_request_without_leaking_provider_message(caplog):
    requests = []

    async def handler(request):
        requests.append(request)
        return httpx.Response(
            400,
            json={"error": {"status": "INVALID_ARGUMENT", "message": "CV privado: api key secret"}},
        )

    provider = GeminiAnalysisProvider(config(), httpx.MockTransport(handler))
    with pytest.raises(AnalisisCvError) as exc:
        await provider.analyze("CV privado", {}, [])
    assert exc.value.status_code == 502
    assert "api_status=INVALID_ARGUMENT" in caplog.text
    assert "CV privado" not in caplog.text
    assert "test-key" not in caplog.text
    assert len(requests) == 1


@pytest.mark.parametrize(
    "body",
    [
        {"candidates": []},
        {"candidates": [{"finishReason": "SAFETY", "content": {"parts": []}}]},
        {"candidates": [{"finishReason": "STOP", "content": {"parts": []}}]},
        {
            "candidates": [
                {
                    "finishReason": "STOP",
                    "content": {"parts": [{"text": '{"unexpected":true}'}]},
                }
            ],
        },
    ],
)
@pytest.mark.asyncio
async def test_provider_invalid_refusal_incomplete(body):
    provider = GeminiAnalysisProvider(
        config(),
        httpx.MockTransport(
            lambda request: httpx.Response(200, json=body),
        ),
    )
    with pytest.raises(AnalisisCvError) as exc:
        await provider.analyze("CV", {}, [])
    assert exc.value.status_code == 502


@pytest.mark.asyncio
async def test_provider_missing_key_timeout_and_http_error():
    without_key = config()
    without_key.gemini_api_key = None
    with pytest.raises(AnalisisCvError, match="GEMINI_API_KEY"):
        await GeminiAnalysisProvider(without_key).analyze("CV", {}, [])

    def timeout(request):
        raise httpx.ReadTimeout("timeout", request=request)

    for handler, status, message in [
        (timeout, 504, "tiempo"),
        (lambda r: httpx.Response(401), 503, "clave API"),
        (lambda r: httpx.Response(429), 503, "limite de uso"),
        (lambda r: httpx.Response(402), 503, "creditos de Gemini"),
        (lambda r: httpx.Response(404), 503, "modelo"),
    ]:
        with pytest.raises(AnalisisCvError) as exc:
            await GeminiAnalysisProvider(config(), httpx.MockTransport(handler)).analyze(
                "CV", {}, []
            )
        assert exc.value.status_code == status
        assert message in str(exc.value)


@pytest.mark.asyncio
@pytest.mark.parametrize("acquisition", [[False], [True, False], [True, True]])
async def test_lock_checks_persistence_and_releases(acquisition, monkeypatch):
    module = importlib.import_module("ssas.analisis_cv.infrastructure.providers.analysis_lock")
    connection = AsyncMock()
    connection.execution_options.return_value = connection
    connection.scalar.side_effect = acquisition + [True]
    engine = MagicMock()

    @asynccontextmanager
    async def dedicated(engine):
        yield connection

    monkeypatch.setattr(module, "_dedicated_connection", dedicated)
    if all(acquisition):
        async with analysis_lock(engine, "id"):
            pass
        assert connection.execute.await_count == 3
    else:
        with pytest.raises(AnalisisCvError) as exc:
            async with analysis_lock(engine, "id"):
                pytest.fail("must not enter")
        assert exc.value.status_code == 409
        assert connection.execute.await_count == (2 if acquisition[0] else 1)


@pytest.mark.asyncio
async def test_read_transaction_never_discards_flushed_writes():
    module = importlib.import_module("ssas.analisis_cv.application.use_cases.analizar_cv")
    session = MagicMock()
    session.new = session.dirty = session.deleted = set()
    session.in_transaction.return_value = True
    session.scalar = AsyncMock(return_value=123)
    session.rollback = AsyncMock()
    with pytest.raises(AnalisisCvError, match="escrituras"):
        await module._release_read_transaction(session)
    session.rollback.assert_not_awaited()
    session.scalar.return_value = None
    await module._release_read_transaction(session)
    session.rollback.assert_awaited_once()


@pytest.mark.asyncio
@pytest.mark.parametrize("failure", [False, True])
async def test_use_case_atomic_persistence_without_commit(monkeypatch, failure):
    module = importlib.import_module("ssas.analisis_cv.application.use_cases.analizar_cv")
    ids = [str(uuid4()) for _ in range(3)]
    row = {
        "postulante_id": str(uuid4()),
        "cv_url": "cv.pdf",
        "experiencia_min": 4,
        "titulo": "Developer",
        "descripcion": "Python",
        "requisitos": "Python",
        "nombres": "Juan",
        "apellidos": "Perez",
        "ci": "123",
        "email": "a@b.org",
        "telefono": "7000000",
        "direccion": None,
        "linkedin": None,
    }
    snapshot = (
        row,
        [{"habilidad_id": "python", "nombre": "Python"}],
        [{"habilidad_id": "python", "peso": 1}],
    )
    monkeypatch.setattr(module, "_snapshot", AsyncMock(return_value=snapshot))
    release = AsyncMock()
    monkeypatch.setattr(module, "_release_read_transaction", release)

    @asynccontextmanager
    async def lock(*args, **kwargs):
        yield

    monkeypatch.setattr(module, "analysis_lock", lock)
    monkeypatch.setattr(module, "lock_persistence", AsyncMock())
    read_session_factory = MagicMock()

    def open_read_session(engine):
        release.assert_awaited_once()
        return read_session_factory(engine)

    monkeypatch.setattr(module, "AsyncSession", open_read_session)
    extractor = SimpleNamespace(
        extract=AsyncMock(return_value="Desarrollo Python. 2 anos de experiencia.")
    )
    monkeypatch.setattr(module, "LocalCvExtractor", lambda settings: extractor)
    provider = SimpleNamespace(analyze=AsyncMock(return_value=result()))
    if failure:
        provider.analyze.side_effect = AnalisisCvError("timeout", 504)
    monkeypatch.setattr(module, "GeminiAnalysisProvider", lambda settings: provider)
    audit = SimpleNamespace(execute=AsyncMock())
    monkeypatch.setattr(module, "RegisterAuditEvent", lambda repository: audit)
    monkeypatch.setattr(module, "SqlAlchemyAuditLogRepository", lambda session: session)
    session = MagicMock()
    session.bind = MagicMock(spec=AsyncEngine)
    session.bind.dialect.name = "postgresql"
    session.new = session.dirty = session.deleted = set()
    session.get = AsyncMock(return_value=SimpleNamespace(puntaje_ia=None))
    session.execute = AsyncMock()
    session.flush = AsyncMock()
    session.commit = AsyncMock()
    session.rollback = AsyncMock()
    if failure:
        with pytest.raises(AnalisisCvError):
            await module.analizar_cv(session, *ids)
        session.add.assert_not_called()
        session.execute.assert_not_awaited()
    else:
        model = await module.analizar_cv(session, *ids, source_ip="127.0.0.1", user_agent="test")
        assert model.puntaje_afinidad == Decimal("90.00")
        assert model.habilidades_detectadas == ["Python"]
        assert model.habilidades_faltantes == []
        from ssas.postulaciones.infrastructure.http.seleccion_schemas import AnalisisResponse

        response = AnalisisResponse.model_validate(model)
        assert response.anios_experiencia_detectados == 2.0
        assert model.modelo_usado == module.settings.gemini_cv_model
        assert response.puntaje_afinidad == Decimal("90.00")
        assert session.get.return_value.puntaje_ia == model.puntaje_afinidad
        sql = str(session.execute.call_args.args[0])
        assert "ON CONFLICT" in sql and "DO NOTHING" in sql
        audit.execute.assert_awaited_once()
        assert audit.execute.call_args.kwargs["source_ip"] == "127.0.0.1"
    release.assert_awaited_once()
    session.commit.assert_not_awaited()
    session.rollback.assert_not_awaited()
