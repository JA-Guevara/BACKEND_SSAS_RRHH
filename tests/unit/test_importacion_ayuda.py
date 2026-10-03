from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from ssas.ayuda.infrastructure.http.router import Question, _rank, _requests, _visible, consultar
from ssas.core.security.hashing import verify_password
from ssas.importacion.application.catalogos import (
    ImportErrorDetail,
    apply,
    parse_csv,
    parse_file,
    preview,
)
from ssas.importacion.application.usuarios import apply_users, preview_users


class FakeSession:
    def __init__(self, by_company):
        self.by_company = by_company
        self.company = None
        self.items = []
        self.flush = AsyncMock()

    async def scalars(self, statement):
        company = statement.compile().params.get("empresa_id_1")
        self.company = company
        return SimpleNamespace(all=lambda: self.by_company.get(company, []))

    def add(self, item):
        self.items.append(item)


def test_csv_rejects_formula_extra_column_and_large_file():
    with pytest.raises(ImportErrorDetail, match="fórmula"):
        parse_csv("habilidades", b"nombre,categoria,descripcion\n=SUM(1),Tecnica,malicioso\n")
    with pytest.raises(ImportErrorDetail, match="adicionales"):
        parse_csv("habilidades", b"nombre,categoria,descripcion\nPython,Tecnica,ok,extra\n")
    with pytest.raises(ImportErrorDetail, match="1 MB"):
        parse_csv("habilidades", b"x" * 1_000_001)


def test_excel_is_read_and_formulas_are_rejected():
    workbook = Workbook()
    workbook.active.append(["nombre", "descripcion"])
    workbook.active.append(["Sin experiencia", "No cumple el perfil"])
    output = BytesIO()
    workbook.save(output)
    _, rows = parse_file("motivos_rechazo", output.getvalue(), "motivos.xlsx")
    assert rows[0]["nombre"] == "Sin experiencia"
    workbook.active["A2"] = "=1+1"
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ImportErrorDetail, match="fórmula"):
        parse_file("motivos_rechazo", output.getvalue(), "motivos.xlsx")


@pytest.mark.asyncio
async def test_department_cycle_and_changed_file_are_rejected():
    session = FakeSession({"company-a": []})
    cycle = b"nombre,descripcion,departamento_padre\nNorte,,Sur\nSur,,Norte\n"
    result = await preview(session, "company-a", "departamentos", cycle)
    assert result["errores"] == 2
    assert session.items == []
    valid = b"nombre,categoria,descripcion\nPython,Tecnica,Lenguaje\n"
    with pytest.raises(ImportErrorDetail, match="cambió"):
        await apply(session, "company-a", "habilidades", valid, "0" * 64)
    assert session.items == []


@pytest.mark.asyncio
async def test_field_length_is_checked_before_write():
    session = FakeSession({"company-a": []})
    content = ("nombre,orden,color,es_inicial,es_contratado,es_rechazado\nEntrevista,2," + "x" * 31 + ",no,no,no\n").encode()
    result = await preview(session, "company-a", "etapas", content)
    assert result["errores"] == 1
    assert session.items == []


@pytest.mark.asyncio
async def test_preview_scopes_existing_names_and_is_read_only():
    existing = SimpleNamespace(nombre="Python")
    session = FakeSession({"company-a": [existing], "company-b": []})
    content = "nombre,categoria,descripcion\nPython,Técnica,Lenguaje\n".encode()
    a = await preview(session, "company-a", "habilidades", content)
    b = await preview(session, "company-b", "habilidades", content)
    assert (a["omitir"], a["crear"]) == (1, 0)
    assert (b["omitir"], b["crear"]) == (0, 1)
    assert session.items == []


@pytest.mark.asyncio
async def test_confirm_repeated_import_omits_existing():
    session = FakeSession({"company-a": []})
    content = b"nombre,categoria,descripcion\nPython,Tecnica,Lenguaje\n"
    digest, _ = parse_csv("habilidades", content)
    first = await apply(session, "company-a", "habilidades", content, digest)
    assert first["crear"] == 1
    assert session.items[0].empresa_id == "company-a"
    session.by_company["company-a"] = session.items
    second = await apply(session, "company-a", "habilidades", content, digest)
    assert second["omitir"] == 1
    assert len(session.items) == 1


@pytest.mark.asyncio
async def test_user_preview_hides_password_and_validates_company_role():
    class UserSession:
        def __init__(self):
            self.roles = [SimpleNamespace(codigo="RECLUTADOR", id="role-a", is_active=True)]

        async def scalars(self, statement):
            items = self.roles if "FROM rol" in str(statement) else []
            return SimpleNamespace(all=lambda: items)

    content = (
        b"nombre,apellido,email,username,contrasena_inicial,rol_codigo,telefono\n"
        b"Ana,Paz,ana@example.com,anapaz,Strong!Pass1234,RECLUTADOR,70000000\n"
    )
    result = await preview_users(UserSession(), "company-a", content, "usuarios.csv")
    assert result["crear"] == 1
    assert "contrasena_inicial" not in result["filas"][0]["datos"]
    assert "Strong!Pass1234" not in str(result)
    wrong_role = content.replace(b"RECLUTADOR", b"ADMIN_AJENO")
    denied = await preview_users(UserSession(), "company-a", wrong_role, "usuarios.csv")
    assert denied["errores"] == 1
    bad_username = content.replace(b"anapaz", b"ana paz")
    invalid = await preview_users(UserSession(), "company-a", bad_username, "usuarios.csv")
    assert invalid["errores"] == 1


@pytest.mark.asyncio
async def test_user_import_hashes_initial_password_and_approves_email(monkeypatch):
    class UserSession:
        async def scalars(self, statement):
            items = [SimpleNamespace(codigo="RECLUTADOR", id="role-a", is_active=True)] if "FROM rol" in str(statement) else []
            return SimpleNamespace(all=lambda: items)

    class UserRepository:
        async def get_by_email(self, *_args):
            return None

        async def get_by_username(self, *_args):
            return None

        async def role_ids_belong_to_empresa(self, roles, company):
            return roles == ["role-a"] and company == "company-a"

        async def create_usuario(self, **kwargs):
            assert kwargs["email_verificado"] is True
            assert kwargs["role_ids"] == ["role-a"]
            assert verify_password("Strong!Pass1234", kwargs["password_hash"])
            return SimpleNamespace(id="new-user")

    monkeypatch.setattr(
        "ssas.importacion.application.usuarios.SqlAlchemyUsuarioRepository",
        lambda _session: UserRepository(),
    )
    content = (
        b"nombre,apellido,email,username,contrasena_inicial,rol_codigo,telefono\n"
        b"Ana,Paz,ana@example.com,anapaz,Strong!Pass1234,RECLUTADOR,70000000\n"
    )
    session = UserSession()
    preview_result = await preview_users(session, "company-a", content, "usuarios.csv")
    result = await apply_users(session, "company-a", content, "usuarios.csv", preview_result["sha256"])
    assert result["crear"] == 1
    assert "Strong!Pass1234" not in str(result)


@pytest.mark.asyncio
async def test_user_import_requires_creation_permission(monkeypatch):
    from ssas.importacion.infrastructure.http.router import authorize_users

    repository = MagicMock()
    repository.get_user_permission_codes = AsyncMock(return_value={"importacion:gestionar"})
    monkeypatch.setattr(
        "ssas.importacion.infrastructure.http.router.SqlAlchemyAuthorizationRepository",
        lambda _session: repository,
    )
    user = SimpleNamespace(id="admin", empresa_id="company-a", es_plataforma=False)
    with pytest.raises(HTTPException) as error:
        await authorize_users("usuarios", user, object())
    assert error.value.status_code == 403


@pytest.mark.asyncio
async def test_excel_template_contains_user_columns_and_no_password_value(monkeypatch):
    from openpyxl import load_workbook

    from ssas.importacion.infrastructure.http.router import plantilla

    monkeypatch.setattr(
        "ssas.importacion.infrastructure.http.router.authorize_users", AsyncMock()
    )
    response = await plantilla("usuarios", "xlsx", SimpleNamespace(), object())
    workbook = load_workbook(BytesIO(response.body), read_only=True)
    assert list(next(workbook.active.values)) == [
        "nombre", "apellido", "email", "username", "contrasena_inicial", "rol_codigo", "telefono"
    ]
    assert workbook.active.max_row == 1
    workbook.close()


@pytest.mark.asyncio
async def test_help_hides_articles_without_permission(monkeypatch):
    repo = MagicMock()
    repo.get_user_permission_codes = AsyncMock(return_value={"postulaciones:ver"})
    monkeypatch.setattr(
        "ssas.ayuda.infrastructure.http.router.SqlAlchemyAuthorizationRepository",
        lambda session: repo,
    )
    user = SimpleNamespace(id="user", empresa_id="company", es_plataforma=False)
    articles = await _visible(user, object())
    assert {item["id"] for item in articles} == {"acceso", "seleccion"}
    assert _rank("contraseña", articles)[0]["id"] == "acceso"
    assert _rank("facturación secreta", articles) == []


@pytest.mark.asyncio
async def test_help_rejects_personal_data_and_rate_limits(monkeypatch):
    monkeypatch.setattr(
        "ssas.ayuda.infrastructure.http.router._visible",
        AsyncMock(return_value=[]),
    )
    user = SimpleNamespace(id="rate-limited-user", empresa_id="company", must_change_password=False)
    _requests.pop(user.id, None)
    with pytest.raises(HTTPException) as personal:
        await consultar(Question(pregunta="Mi correo es persona@example.com"), user, object())
    assert personal.value.status_code == 422
    for _ in range(9):
        result = await consultar(Question(pregunta="Pregunta desconocida"), user, object())
        assert result["modo"] == "sin_resultado"
    with pytest.raises(HTTPException) as limited:
        await consultar(Question(pregunta="Pregunta desconocida"), user, object())
    assert limited.value.status_code == 429
    _requests.pop(user.id, None)
