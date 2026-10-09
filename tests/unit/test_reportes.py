import inspect
import re
from datetime import UTC, datetime
from decimal import Decimal
from email.message import EmailMessage
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from ssas.auth.infrastructure.persistence.models.user import UserModel
from ssas.postulaciones.infrastructure.persistence.models.postulacion import PostulacionModel
from ssas.postulantes.infrastructure.persistence.models.postulante import PostulanteModel
from ssas.reportes.application.cache import CacheAgregados, cache_agregados
from ssas.reportes.application.cuotas import (
    CUOTAS,
    cuota_de_empresa,
    cuota_de_plan,
    exportaciones_de_hoy,
    interpretaciones_ia_de_hoy,
    verificar_exportaciones_dia,
    verificar_filas,
    verificar_interpretaciones_ia,
)
from ssas.reportes.domain.catalogo import (
    CATALOGO,
    Campo,
    Fuente,
    Sensibilidad,
    TipoCampo,
)
from ssas.reportes.infrastructure.http import router as router_module
from ssas.reportes.infrastructure.http.router import (
    FROM_SQL,
    LIMITE_FILAS,
    SOURCES,
    TENANT_COLUMN,
    _autorizar,
    _autorizar_agregado,
    _campo_visible,
    _catalogo_visible,
    _columnas_sensibles,
    _contar,
    _document,
    _rows,
    _rows_truncado,
    _serialize_ejecucion,
    _validate,
)
from ssas.reportes.infrastructure.http.schemas import ReporteConfig
from ssas.reportes.infrastructure.http.schemas_agregado import (
    Agregacion,
    ConsultaAgregada,
    Medida,
    RespuestaAgregada,
)
from ssas.reportes.infrastructure.persistence.models.reporte import ReporteEjecucionModel
from ssas.vacantes.infrastructure.persistence.models.vacante import VacanteModel


def test_report_catalog_rejects_unknown_columns() -> None:
    config = ReporteConfig(fuente="usuarios", columnas=["password_hash"])

    with pytest.raises(HTTPException) as error:
        _validate(config)

    assert error.value.status_code == 422


@pytest.mark.parametrize(
    ("format", "signature"),
    [("html", b"<!doctype html>"), ("xlsx", b"PK"), ("pdf", b"%PDF"), ("csv", b"\xef\xbb\xbf")],
)
def test_exports_create_real_documents(format: str, signature: bytes) -> None:
    config = ReporteConfig(fuente="usuarios", columnas=["nombres", "email"])

    content, media_type = _document(
        config, [{"nombres": "Ana", "email": "ana@example.com"}], format
    )

    assert content.startswith(signature)
    assert media_type



def test_xlsx_tiene_encabezado_inmovilizado_y_hoja_de_parametros() -> None:
    config = ReporteConfig(fuente="usuarios", columnas=["nombres", "email"])

    content, _ = _document(
        config,
        [{"nombres": "Ana", "email": "ana@example.com"}],
        "xlsx",
        empresa="ACME",
    )

    book = load_workbook(BytesIO(content))
    assert book["Datos"].freeze_panes == "A2"
    assert "Parámetros" in book.sheetnames
    resumen = {fila[0].value: fila[1].value for fila in book["Parámetros"].iter_rows()}
    assert resumen["Empresa"] == "ACME"
    assert resumen["Registros"] == 1


def test_xlsx_convierte_decimales_y_fechas_con_zona() -> None:
    config = ReporteConfig(fuente="usuarios", columnas=["nombres", "email"])
    fila = {
        "nombres": Decimal("10.50"),
        "email": datetime(2026, 10, 8, 12, 0, tzinfo=UTC),
    }

    content, _ = _document(config, [fila], "xlsx")

    book = load_workbook(BytesIO(content))
    hoja = book["Datos"]
    assert isinstance(hoja["A2"].value, float)
    assert hoja["B2"].value.tzinfo is None


@pytest.mark.asyncio
async def test_postulaciones_report_uses_real_columns_and_tenant_scope() -> None:
    session = MagicMock()
    result = MagicMock()
    result.mappings.return_value.all.return_value = []
    session.execute = AsyncMock(return_value=result)
    config = ReporteConfig(
        fuente="postulaciones",
        columnas=["postulante", "puntaje"],
        filtros=[{"campo": "estado", "operador": "igual", "valor": "ACTIVA"}],
    )

    assert await _rows(session, "empresa-a", config) == []
    statement, params = session.execute.await_args.args
    sql = str(statement)
    assert "v.empresa_id = :empresa_id" in sql
    assert "p.empresa_id=v.empresa_id" in sql
    assert "COALESCE(po.puntaje_manual, po.puntaje_ia) AS puntaje" in sql
    assert "po.empresa_id" not in sql
    assert "po.puntaje_final" not in sql
    assert params == {"empresa_id": "empresa-a", "v0": "ACTIVA"}


def test_all_report_columns_exist_in_the_database_models() -> None:
    aliases = {
        "v": set(VacanteModel.__table__.columns.keys()),
        "u": set(UserModel.__table__.columns.keys()),
        "po": set(PostulacionModel.__table__.columns.keys()),
        "p": set(PostulanteModel.__table__.columns.keys()),
    }
    for source, columns in SOURCES.items():
        expressions = [*columns.values(), FROM_SQL[source], TENANT_COLUMN[source]]
        for expression in expressions:
            for alias, column in re.findall(r"\b(v|u|po|p)\.([a-z_]+)\b", expression):
                assert column in aliases[alias], f"{source}: {alias}.{column} no existe"


@pytest.mark.asyncio
@pytest.mark.parametrize("source, field", [
    ("vacantes", "estado"),
    ("usuarios", "username"),
    ("postulaciones", "postulante"),
])
@pytest.mark.parametrize("operator, value, fragment", [
    ("igual", "ACTIVA", " = :v0"),
    ("contiene", "Ana", " ILIKE :v0"),
    ("mayor_igual", "A", " >= :v0"),
    ("menor_igual", "Z", " <= :v0"),
])
async def test_filters_work_for_every_report_source(
    source: str, field: str, operator: str, value: str, fragment: str
) -> None:
    session = MagicMock()
    result = MagicMock()
    result.mappings.return_value.all.return_value = []
    session.execute = AsyncMock(return_value=result)
    config = ReporteConfig(
        fuente=source,
        columnas=list(SOURCES[source]),
        filtros=[{"campo": field, "operador": operator, "valor": value}],
        orden=[{"campo": field, "direccion": "desc"}],
    )

    assert await _rows(session, "empresa-a", config) == []
    statement, params = session.execute.await_args.args
    sql = str(statement)
    assert f"{TENANT_COLUMN[source]} = :empresa_id" in sql
    assert fragment in sql
    assert " ORDER BY " in sql and " DESC" in sql
    assert params["empresa_id"] == "empresa-a"
    assert params["v0"] == (f"%{value}%" if operator == "contiene" else value)


@pytest.mark.asyncio
@pytest.mark.parametrize("source, date_field", [
    ("vacantes", "fecha_publicacion"),
    ("usuarios", "ultimo_acceso"),
    ("postulaciones", "fecha_postulacion"),
])
async def test_date_range_filter_works_for_every_report_source(
    source: str, date_field: str
) -> None:
    session = MagicMock()
    result = MagicMock()
    result.mappings.return_value.all.return_value = []
    session.execute = AsyncMock(return_value=result)
    config = ReporteConfig(
        fuente=source,
        columnas=[date_field],
        filtros=[{
            "campo": date_field,
            "operador": "entre",
            "valor": ["2026-10-01", "2026-10-31"],
        }],
    )

    assert await _rows(session, "empresa-a", config) == []
    statement, params = session.execute.await_args.args
    assert " BETWEEN :v0a AND :v0b" in str(statement)
    assert params == {
        "empresa_id": "empresa-a",
        "v0a": "2026-10-01",
        "v0b": "2026-10-31",
    }


def test_catalogo_oculta_fuentes_sin_permiso() -> None:
    assert _catalogo_visible(set(), False) == []
    codigos = {fuente["codigo"] for fuente in _catalogo_visible({"vacantes:ver"}, False)}
    assert codigos == {"vacantes"}


def test_catalogo_de_plataforma_ve_todas_las_fuentes() -> None:
    codigos = {fuente["codigo"] for fuente in _catalogo_visible(set(), True)}
    assert codigos == set(CATALOGO)


def test_campo_con_permiso_oculto_sin_autorizacion() -> None:
    campo = Campo(
        codigo="salario", etiqueta="Salario", sql="e.salario",
        tipo=TipoCampo.NUMERO, permiso="nomina:ver",
    )

    assert _campo_visible(campo, set(), False) is False
    assert _campo_visible(campo, {"nomina:ver"}, False) is True
    assert _campo_visible(campo, set(), True) is True


@pytest.mark.asyncio
async def test_truncado_se_informa_cuando_hay_mas_filas_que_el_limite() -> None:
    session = MagicMock()
    result = MagicMock()
    result.mappings.return_value.all.return_value = [
        {"nombres": f"p{i}"} for i in range(LIMITE_FILAS + 1)
    ]
    session.execute = AsyncMock(return_value=result)
    config = ReporteConfig(fuente="usuarios", columnas=["nombres"])

    rows, truncado = await _rows_truncado(session, "empresa-a", config)

    assert truncado is True
    assert len(rows) == LIMITE_FILAS
    assert f"LIMIT {LIMITE_FILAS + 1}" in str(session.execute.await_args.args[0])


@pytest.mark.asyncio
async def test_conteo_cuenta_con_los_mismos_filtros() -> None:
    session = MagicMock()
    session.scalar = AsyncMock(return_value=12_345)
    session.execute = AsyncMock()
    config = ReporteConfig(
        fuente="postulaciones",
        columnas=["postulante"],
        filtros=[{"campo": "estado", "operador": "igual", "valor": "ACTIVA"}],
    )

    total = await _contar(session, "empresa-a", config)

    assert total == 12_345
    sql = str(session.scalar.await_args.args[0])
    assert "count(*) AS total" in sql
    assert "v.empresa_id = :empresa_id" in sql


def _registrar_fuente_prueba(monkeypatch, fuente: Fuente) -> None:
    monkeypatch.setitem(CATALOGO, fuente.codigo, fuente)
    monkeypatch.setitem(
        SOURCES, fuente.codigo, {name: campo.sql for name, campo in fuente.campos.items()}
    )
    monkeypatch.setitem(FROM_SQL, fuente.codigo, fuente.from_sql)
    monkeypatch.setitem(TENANT_COLUMN, fuente.codigo, fuente.columna_tenant)


def _fuente_de_prueba(codigo: str) -> Fuente:
    return Fuente(
        codigo=codigo,
        etiqueta="Fuente de prueba",
        descripcion="",
        from_sql="tabla t",
        columna_tenant="t.empresa_id",
        permiso="test:ver",
        campos={
            "salario": Campo(
                codigo="salario",
                etiqueta="Salario",
                sql="t.salario",
                tipo=TipoCampo.NUMERO,
                permiso="nomina:ver",
            ),
            "email": Campo(
                codigo="email",
                etiqueta="Correo",
                sql="t.email",
                tipo=TipoCampo.TEXTO,
                sensibilidad=Sensibilidad.PERSONAL,
            ),
            "nota": Campo(
                codigo="nota",
                etiqueta="Nota",
                sql="t.nota",
                tipo=TipoCampo.TEXTO,
                sensibilidad=Sensibilidad.CONFIDENCIAL,
            ),
            "nombre": Campo(
                codigo="nombre",
                etiqueta="Nombre",
                sql="t.nombre",
                tipo=TipoCampo.TEXTO,
            ),
        },
    )


def test_campo_con_permiso_no_autorizado_es_403(monkeypatch) -> None:
    _registrar_fuente_prueba(monkeypatch, _fuente_de_prueba("test_perm"))
    config = ReporteConfig(fuente="test_perm", columnas=["salario"])

    with pytest.raises(HTTPException) as error:
        _autorizar(config, set(), False)
    assert error.value.status_code == 403

    assert _autorizar(config, {"nomina:ver"}, False)["salario"] == "t.salario"
    assert _autorizar(config, set(), True)["salario"] == "t.salario"


def test_permiso_se_exige_tambien_en_filtros(monkeypatch) -> None:
    _registrar_fuente_prueba(monkeypatch, _fuente_de_prueba("test_perm"))
    config = ReporteConfig(
        fuente="test_perm",
        columnas=["nombre"],
        filtros=[{"campo": "salario", "operador": "mayor_igual", "valor": 1}],
    )

    with pytest.raises(HTTPException) as error:
        _autorizar(config, set(), False)
    assert error.value.status_code == 403


def test_columnas_sensibles_marca_personal_y_confidencial(monkeypatch) -> None:
    _registrar_fuente_prueba(monkeypatch, _fuente_de_prueba("test_sens"))
    config = ReporteConfig(
        fuente="test_sens", columnas=["email", "nombre", "nota"]
    )

    assert _columnas_sensibles(config) == ["email", "nota"]


def test_agregado_rechaza_campo_sin_permiso(monkeypatch) -> None:
    _registrar_fuente_prueba(monkeypatch, _fuente_de_prueba("test_perm"))
    consulta = ConsultaAgregada(
        fuente="test_perm",
        medidas=[Medida(agregacion=Agregacion.SUMA, campo="salario")],
    )

    with pytest.raises(HTTPException) as error:
        _autorizar_agregado(consulta, set(), False)
    assert error.value.status_code == 403

    _autorizar_agregado(consulta, {"nomina:ver"}, False)


def test_envio_smtp_usa_timeout_y_tls_fuera_del_bucle(monkeypatch) -> None:
    registro: dict[str, object] = {}

    class FakeSMTP:
        def __init__(self, host, port, timeout=None):
            registro["host"] = host
            registro["port"] = port
            registro["timeout"] = timeout

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def starttls(self):
            registro["tls"] = True

        def login(self, username, password):
            registro["login"] = (username, password)

        def send_message(self, message):
            registro["message"] = message

    monkeypatch.setattr(router_module, "smtplib", SimpleNamespace(SMTP=FakeSMTP))
    monkeypatch.setattr(router_module.settings, "smtp_host", "smtp.test")
    monkeypatch.setattr(router_module.settings, "smtp_port", 2525)
    monkeypatch.setattr(router_module.settings, "smtp_use_tls", True)
    monkeypatch.setattr(router_module.settings, "smtp_username", "usuario")
    monkeypatch.setattr(router_module.settings, "smtp_password", "clave")
    message = EmailMessage()

    router_module._enviar_correo(message)

    assert registro["host"] == "smtp.test"
    assert registro["port"] == 2525
    assert registro["timeout"] == router_module.settings.smtp_timeout_seconds
    assert registro["tls"] is True
    assert registro["login"] == ("usuario", "clave")
    assert registro["message"] is message


def test_el_envio_se_delega_a_un_hilo() -> None:
    source = inspect.getsource(router_module.enviar)

    assert "anyio.to_thread.run_sync(_enviar_correo" in source
    assert "with smtplib.SMTP" not in source


def test_ejecucion_serializa_estado_y_error() -> None:
    item = ReporteEjecucionModel(
        id="e1",
        empresa_id="emp1",
        usuario_id="u1",
        formato="xlsx",
        estado="ERROR",
        error="RuntimeError: boom",
        filtros_aplicados=[],
        columnas_sensibles=[],
        fecha_inicio=datetime.now(UTC),
        fecha_fin=None,
    )

    data = _serialize_ejecucion(item)

    assert data.estado == "ERROR"
    assert data.error == "RuntimeError: boom"
    assert data.cantidad_registros is None


@pytest.mark.asyncio
async def test_exportar_fallido_registra_estado_error(monkeypatch) -> None:
    monkeypatch.setattr(router_module, "_permisos", AsyncMock(return_value=set()))

    async def boom(*args, **kwargs):
        raise RuntimeError("db caida")

    monkeypatch.setattr(router_module, "_rows", boom)
    session = MagicMock()
    session.scalar = AsyncMock(return_value=0)
    session.execute = AsyncMock()
    session.commit = AsyncMock()
    user = SimpleNamespace(id="u1", empresa_id="emp1", es_plataforma=False)
    body = ReporteConfig(fuente="usuarios", columnas=["nombres"])

    with pytest.raises(HTTPException) as error:
        await router_module.exportar("xlsx", body, None, None, user, session)

    assert error.value.status_code == 500
    ejecucion = session.add.call_args.args[0]
    assert ejecucion.estado == "ERROR"
    assert ejecucion.error.startswith("RuntimeError: db caida")
    session.commit.assert_awaited_once()


def test_cuota_excedida_es_429() -> None:
    basico = cuota_de_plan("Básico")

    with pytest.raises(HTTPException) as error:
        verificar_exportaciones_dia(basico, 21)

    assert error.value.status_code == 429
    assert "Básico" in error.value.detail
    assert "20" in error.value.detail

    verificar_exportaciones_dia(basico, 19)


def test_filas_sobre_el_limite_del_plan_es_429() -> None:
    basico = cuota_de_plan("Basico")

    with pytest.raises(HTTPException) as error:
        verificar_filas(basico, 1_001)

    assert error.value.status_code == 429
    assert "1000" in error.value.detail
    assert "Básico" in error.value.detail

    verificar_filas(basico, 1_000)


def test_cuota_de_ia_excedida_es_429() -> None:
    with pytest.raises(HTTPException) as error:
        verificar_interpretaciones_ia(cuota_de_plan("basico"), 10)

    assert error.value.status_code == 429
    assert "IA" in error.value.detail


def test_plan_premium_no_tiene_limite_diario() -> None:
    premium = cuota_de_plan("Empresarial")

    assert premium is CUOTAS["premium"]
    verificar_exportaciones_dia(premium, 10_000)


def test_plan_desconocido_cae_al_mas_restrictivo() -> None:
    assert cuota_de_plan(None) is CUOTAS["basico"]
    assert cuota_de_plan("Plan Raro") is CUOTAS["basico"]


@pytest.mark.asyncio
async def test_cuota_de_empresa_usa_el_plan_activo() -> None:
    session = MagicMock()
    session.scalar = AsyncMock(return_value="Profesional")

    assert await cuota_de_empresa(session, "emp1") is CUOTAS["profesional"]
    sql = str(session.scalar.await_args.args[0])
    assert "plan_suscripcion" in sql
    assert "suscripcion.estado = :estado_1" in sql


@pytest.mark.asyncio
async def test_conteos_diarios_leen_reporte_ejecucion() -> None:
    session = MagicMock()
    session.scalar = AsyncMock(return_value=7)

    assert await exportaciones_de_hoy(session, "emp1") == 7
    assert await interpretaciones_ia_de_hoy(session, "emp1") == 7
    sql = str(session.scalar.await_args.args[0])
    assert "reporte_ejecucion.formato IN" in sql


def test_cache_no_cruza_empresas() -> None:
    cache = CacheAgregados()
    cache.set("empresa-a", "misma-consulta", ["dato-a"])

    assert cache.get("empresa-a", "misma-consulta") == ["dato-a"]
    assert cache.get("empresa-b", "misma-consulta") is None


def test_cache_expira_a_los_60_segundos() -> None:
    ahora = [0.0]
    cache = CacheAgregados(ttl_segundos=60, reloj=lambda: ahora[0])
    cache.set("emp1", "clave", "valor")

    assert cache.get("emp1", "clave") == "valor"
    ahora[0] = 59.9
    assert cache.get("emp1", "clave") == "valor"
    ahora[0] = 60.0
    assert cache.get("emp1", "clave") is None


@pytest.mark.asyncio
async def test_agregado_reutiliza_la_cache(monkeypatch) -> None:
    monkeypatch.setattr(router_module, "_permisos", AsyncMock(return_value=set()))
    llamadas: list[str] = []

    async def fake_ejecutar(session, consulta, empresa_id):
        llamadas.append(empresa_id)
        return RespuestaAgregada(
            series=[],
            medidas=["conteo"],
            total_grupos=0,
            truncado=False,
            generado_en=datetime.now(UTC),
            milisegundos=1,
        )

    monkeypatch.setattr(router_module, "ejecutar_agregado", fake_ejecutar)
    cache_agregados.clear()
    session = MagicMock()
    session.scalar = AsyncMock(return_value=None)
    user = SimpleNamespace(id="u1", empresa_id="emp1", es_plataforma=False)
    body = ConsultaAgregada(fuente="postulaciones", medidas=[{"agregacion": "conteo"}])

    primera = await router_module.agregado(body, None, user, session)
    segunda = await router_module.agregado(body, None, user, session)

    assert llamadas == ["emp1"]
    assert primera is segunda


@pytest.mark.asyncio
async def test_statement_timeout_se_aplica_antes_de_consultar() -> None:
    session = MagicMock()
    result = MagicMock()
    result.mappings.return_value.all.return_value = []
    session.execute = AsyncMock(return_value=result)
    config = ReporteConfig(fuente="usuarios", columnas=["nombres"])

    await _rows(session, "empresa-a", config)

    sentencias = [str(call.args[0]) for call in session.execute.await_args_list]
    assert sentencias[0] == "SET LOCAL statement_timeout = '15s'"
    assert "SELECT" in sentencias[1]


def test_csv_tiene_bom_y_etiquetas_en_espanol() -> None:
    config = ReporteConfig(fuente="usuarios", columnas=["nombres", "email"])
    fila = {"nombres": "José Pérez", "email": "jose@example.com"}

    content, media_type = _document(config, [fila], "csv")

    assert content.startswith(b"\xef\xbb\xbf")
    assert media_type == "text/csv; charset=utf-8"
    texto = content.decode("utf-8-sig")
    assert "Nombres,Correo" in texto
    assert "José Pérez,jose@example.com" in texto


def test_pdf_platypus_genera_documento_valido_con_metadatos() -> None:
    config = ReporteConfig(fuente="postulaciones", columnas=["postulante", "estado"])
    fila = {"postulante": "Carlos Gómez", "estado": "ACTIVA"}

    content, media_type = _document(
        config,
        [fila],
        "pdf",
        empresa="Empresa Demo",
        generado_en=datetime(2026, 10, 8, 15, 0, tzinfo=UTC),
    )

    assert content.startswith(b"%PDF")
    assert media_type == "application/pdf"
    assert len(content) > 1000


@pytest.mark.asyncio
async def test_exportar_rechaza_formato_html(monkeypatch) -> None:
    monkeypatch.setattr(router_module, "_permisos", AsyncMock(return_value=set()))
    session = MagicMock()
    user = SimpleNamespace(id="u1", empresa_id="emp1", es_plataforma=False)
    request = MagicMock()
    config = ReporteConfig(fuente="usuarios", columnas=["nombres"])

    with pytest.raises(HTTPException) as error:
        await router_module.exportar("html", config, request, None, user, session)

    assert error.value.status_code == 422
    assert "html" in error.value.detail


@pytest.mark.asyncio
async def test_panel_devuelve_widgets_por_defecto_si_no_hay_guardados(monkeypatch) -> None:
    monkeypatch.setattr(router_module, "_permisos", AsyncMock(return_value={"reportes:ver", "postulaciones:ver", "vacantes:ver", "analisis_cv:ver"}))
    session = MagicMock()
    result = MagicMock()
    result.scalars.return_value.all.return_value = []
    session.execute = AsyncMock(return_value=result)
    user = SimpleNamespace(id="u1", empresa_id="emp1", es_plataforma=False)

    respuesta = await router_module.obtener_panel(None, user, session)

    assert respuesta.origen == "predeterminadas"
    assert len(respuesta.widgets) >= 6
    titulos = [w.titulo for w in respuesta.widgets]
    assert "Postulaciones" in titulos
    assert "Embudo de selección" in titulos
    assert "Postulaciones por semana" in titulos
    assert respuesta.omitidas_por_permiso == []
    assert "postulaciones" in respuesta.fuentes_disponibles


@pytest.mark.asyncio
async def test_panel_sin_permisos_explica_lo_omitido(monkeypatch) -> None:
    """Un admin con solo reportes:ver nunca ve el panel mudo: explica qué falta."""
    monkeypatch.setattr(router_module, "_permisos", AsyncMock(return_value={"reportes:ver"}))
    session = MagicMock()
    result = MagicMock()
    result.scalars.return_value.all.return_value = []
    session.execute = AsyncMock(return_value=result)
    user = SimpleNamespace(id="u1", empresa_id="emp1", es_plataforma=False)

    respuesta = await router_module.obtener_panel(None, user, session)

    # La tarjeta universal (solo reportes:ver) siempre está; el resto se explica.
    assert [w.titulo for w in respuesta.widgets] == ["Ejecuciones de reportes hoy"]
    assert any(s.startswith("Postulaciones — requiere «postulaciones:ver»") for s in respuesta.omitidas_por_permiso)
    assert any(s.startswith("Vacantes activas — requiere «vacantes:ver»") for s in respuesta.omitidas_por_permiso)
    assert respuesta.fuentes_disponibles == ["reportes"]


@pytest.mark.asyncio
async def test_panel_fijar_tarjeta_excede_cuota_es_429(monkeypatch) -> None:
    monkeypatch.setattr(router_module, "_permisos", AsyncMock(return_value={"postulaciones:ver"}))
    monkeypatch.setattr(router_module, "tarjetas_fijadas_de_usuario", AsyncMock(return_value=4))
    session = MagicMock()
    session.scalar = AsyncMock(return_value=None)  # Plan básico -> 4 tarjetas max
    user = SimpleNamespace(id="u1", empresa_id="emp1", es_plataforma=False)
    body = router_module.CrearWidgetPanel(
        titulo="Nueva tarjeta",
        tipo="kpi",
        consulta=ConsultaAgregada(fuente="postulaciones", medidas=[{"agregacion": "conteo"}]),
    )

    with pytest.raises(HTTPException) as error:
        await router_module.fijar_tarjeta(body, None, user, session)

    assert error.value.status_code == 429
    assert "tarjetas" in error.value.detail

