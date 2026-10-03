"""CSV/Excel catalog imports scoped to exactly one company."""

import csv
import hashlib
import io
from datetime import date, datetime
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from sqlalchemy import select

from ssas.cargos.infrastructure.persistence.models.cargo import CargoModel
from ssas.departamentos.infrastructure.persistence.models.departamento import DepartamentoModel
from ssas.habilidades.infrastructure.persistence.models.habilidad import HabilidadModel
from ssas.postulaciones.infrastructure.persistence.models.etapa_reclutamiento import (
    EtapaReclutamientoModel,
)
from ssas.postulaciones.infrastructure.persistence.models.motivo_rechazo import MotivoRechazoModel

CATALOGS = {
    "departamentos": (DepartamentoModel, ("nombre", "descripcion", "departamento_padre")),
    "cargos": (CargoModel, ("nombre", "departamento", "descripcion", "nivel")),
    "habilidades": (HabilidadModel, ("nombre", "categoria", "descripcion")),
    "etapas": (EtapaReclutamientoModel, ("nombre", "orden", "color", "es_inicial", "es_contratado", "es_rechazado")),
    "motivos_rechazo": (MotivoRechazoModel, ("nombre", "descripcion")),
    "usuarios": (None, ("nombre", "apellido", "email", "username", "contrasena_inicial", "rol_codigo", "telefono")),
    "empleados": (None, (
        "codigo", "nombres", "apellido_paterno", "apellido_materno", "ci", "ci_expedido",
        "fecha_ingreso", "estado", "fecha_nacimiento", "genero", "estado_civil", "direccion",
        "telefono", "email_personal", "contacto_emergencia", "telefono_emergencia",
        "nua_cua", "afp", "banco", "numero_cuenta", "tipo_cuenta", "fecha_salida", "motivo_salida",
    )),
}
STAGE_FLAGS = ("es_inicial", "es_contratado", "es_rechazado")


class ImportErrorDetail(ValueError):
    pass


def parse_file(kind: str, content: bytes, filename: str = "datos.csv"):
    if kind not in CATALOGS:
        raise ImportErrorDetail("Catálogo no admitido")
    if not content or len(content) > 1_000_000:
        raise ImportErrorDetail("El archivo debe ocupar entre 1 byte y 1 MB")
    try:
        expected = CATALOGS[kind][1]
        max_rows = 100 if kind == "usuarios" else 1000
        if filename.lower().endswith(".xlsx"):
            with ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 100 or sum(i.file_size for i in archive.infolist()) > 10_000_000:
                    raise ImportErrorDetail("El Excel supera el límite de tamaño interno")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.sheetnames) != 1:
                    raise ImportErrorDetail("El Excel debe contener una sola hoja")
                iterator = workbook.active.iter_rows(values_only=True)
                header = next(iterator, ())
                if list(header) != list(expected):
                    raise ImportErrorDetail("Encabezados esperados: " + ", ".join(expected))
                raw_rows = []
                for cells in iterator:
                    if len(raw_rows) >= max_rows:
                        raise ImportErrorDetail(f"El Excel supera {max_rows} filas")
                    if len(cells) > len(expected) and any(c is not None for c in cells[len(expected):]):
                        raise ImportErrorDetail("Una fila tiene columnas adicionales")
                    if any(c is not None for c in cells):
                        if kind == "empleados" and any(
                            cells[index] is not None and not isinstance(cells[index], str)
                            for index, column in enumerate(expected)
                            if column in {"codigo", "ci", "ci_expedido", "telefono", "numero_cuenta"}
                            and index < len(cells)
                        ):
                            raise ImportErrorDetail("Código, CI, teléfono y cuenta deben tener formato Texto en Excel")
                        raw_rows.append(dict(zip(expected, cells, strict=False)))
            finally:
                workbook.close()
        elif filename.lower().endswith(".csv"):
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
            if reader.fieldnames != list(expected):
                raise ImportErrorDetail("Encabezados esperados: " + ", ".join(expected))
            raw_rows = []
            for row in reader:
                if len(raw_rows) >= max_rows:
                    raise ImportErrorDetail(f"El CSV supera {max_rows} filas")
                if None in row:
                    raise ImportErrorDetail("Una fila tiene columnas adicionales")
                raw_rows.append(row)
        else:
            raise ImportErrorDetail("Selecciona un archivo .csv o .xlsx")
        rows = []
        for row in raw_rows:
            clean = {
                key: (
                    value.date().isoformat() if isinstance(value, datetime) and kind == "empleados" and key.startswith("fecha_")
                    else value.isoformat() if isinstance(value, date) and kind == "empleados" and key.startswith("fecha_")
                    else str(value) if key == "contrasena_inicial" else str(value).strip()
                )
                if value is not None else "" for key, value in row.items()
            }
            if any(value.lstrip().startswith(("=", "+", "-", "@")) for value in clean.values()):
                raise ImportErrorDetail("El archivo contiene una fórmula o celda no permitida")
            rows.append(clean)
    except ImportErrorDetail:
        raise
    except (UnicodeDecodeError, csv.Error, BadZipFile, ValueError, TypeError) as exc:
        raise ImportErrorDetail("El archivo CSV/Excel no es válido") from exc
    if not rows:
        raise ImportErrorDetail("El archivo no contiene filas")
    return hashlib.sha256(content).hexdigest(), rows


def parse_csv(kind: str, content: bytes):
    return parse_file(kind, content)


async def preview(session, company_id: str, kind: str, content: bytes, filename: str = "datos.csv"):
    digest, rows = parse_file(kind, content, filename)
    model = CATALOGS[kind][0]
    existing = (await session.scalars(select(model).where(model.empresa_id == company_id))).all()
    names = {item.nombre.casefold(): item for item in existing}
    departments = (
        await session.scalars(
            select(DepartamentoModel).where(DepartamentoModel.empresa_id == company_id)
        )
    ).all()
    department_names = {item.nombre.casefold() for item in departments}
    source_names = {r["nombre"].casefold() for r in rows if r["nombre"]}
    edges = {r["nombre"].casefold(): r.get("departamento_padre", "").casefold() for r in rows}
    seen: set[str] = set()
    orders: set[int] = set()
    used_orders = {item.orden for item in existing} if kind == "etapas" else set()
    existing_flags = {flag for flag in STAGE_FLAGS if kind == "etapas" and any(getattr(item, flag) for item in existing)}
    seen_flags: set[str] = set()
    result = []
    for number, row in enumerate(rows, 2):
        name = row["nombre"]
        problems = []
        if len(name) < 2 or len(name) > (100 if kind == "etapas" else 120):
            problems.append("Nombre inválido")
        if name.casefold() in seen:
            problems.append("Nombre repetido en el archivo")
        seen.add(name.casefold())
        limits = {"descripcion": 4000, "nivel": 80, "color": 30}
        if any(len(value) > limits.get(field, 120) for field, value in row.items()):
            problems.append("Un campo excede su longitud máxima")
        if kind == "departamentos" and row["departamento_padre"]:
            parent = row["departamento_padre"].casefold()
            if parent not in source_names and parent not in department_names:
                problems.append("Departamento padre inexistente")
            visited, current = set(), name.casefold()
            while edges.get(current):
                if current in visited:
                    problems.append("Jerarquía circular")
                    break
                visited.add(current)
                current = edges[current]
        if kind == "cargos" and row["departamento"].casefold() not in department_names:
            problems.append("Departamento inexistente en esta empresa")
        if kind == "etapas":
            try:
                order = int(row["orden"])
                if not 1 <= order <= 1000 or order in orders:
                    raise ValueError()
                orders.add(order)
                if order in used_orders and name.casefold() not in names:
                    problems.append("Orden ya ocupado")
            except ValueError:
                problems.append("Orden inválido o repetido")
            for flag in STAGE_FLAGS:
                value = row[flag].casefold()
                if value not in {"si", "no", "true", "false", "1", "0"}:
                    problems.append(f"{flag}: usa si o no")
                elif value in {"si", "true", "1"}:
                    if flag in seen_flags or (flag in existing_flags and name.casefold() not in names):
                        problems.append(f"{flag} ya está asignado")
                    seen_flags.add(flag)
        action = "error" if problems else ("omitir" if name.casefold() in names else "crear")
        result.append({"fila": number, "datos": row, "accion": action, "errores": problems})
    if kind == "etapas" and "es_inicial" not in existing_flags and "es_inicial" not in seen_flags:
        result[0]["errores"].append("Debes indicar una etapa inicial")
        result[0]["accion"] = "error"
    return {
        "sha256": digest,
        "filas": result,
        "crear": sum(r["accion"] == "crear" for r in result),
        "omitir": sum(r["accion"] == "omitir" for r in result),
        "errores": sum(r["accion"] == "error" for r in result),
    }


async def apply(session, company_id: str, kind: str, content: bytes, expected_digest: str, filename: str = "datos.csv"):
    result = await preview(session, company_id, kind, content, filename)
    if result["sha256"] != expected_digest:
        raise ImportErrorDetail("El archivo cambió desde la vista previa")
    if result["errores"]:
        raise ImportErrorDetail("Corrige las filas marcadas antes de importar")
    existing = (
        await session.scalars(
            select(DepartamentoModel).where(DepartamentoModel.empresa_id == company_id)
        )
    ).all()
    departments = {item.nombre.casefold(): item.id for item in existing}
    pending = [r["datos"] for r in result["filas"] if r["accion"] == "crear"]
    while pending:
        later = []
        inserted = 0
        for row in pending:
            data = {"id": str(uuid4()), "empresa_id": company_id, "nombre": row["nombre"]}
            if kind == "departamentos":
                parent = row["departamento_padre"].casefold()
                if parent and parent not in departments:
                    later.append(row)
                    continue
                data.update(
                    descripcion=row["descripcion"] or None,
                    departamento_padre_id=departments.get(parent),
                )
                item = DepartamentoModel(**data)
                departments[row["nombre"].casefold()] = data["id"]
            elif kind == "cargos":
                data.update(
                    departamento_id=departments[row["departamento"].casefold()],
                    descripcion=row["descripcion"] or None,
                    nivel=row["nivel"] or None,
                )
                item = CargoModel(**data)
            elif kind == "habilidades":
                data.update(
                    categoria=row["categoria"] or None, descripcion=row["descripcion"] or None
                )
                item = HabilidadModel(**data)
            elif kind == "motivos_rechazo":
                data.update(descripcion=row["descripcion"] or None)
                item = MotivoRechazoModel(**data)
            else:
                data.update(
                    orden=int(row["orden"]), color=row["color"] or None,
                    **{flag: row[flag].casefold() in {"si", "true", "1"} for flag in STAGE_FLAGS},
                )
                item = EtapaReclutamientoModel(**data)
            session.add(item)
            inserted += 1
        if not inserted:
            raise ImportErrorDetail("No se pudo resolver la jerarquía de departamentos")
        await session.flush()
        pending = later
    return result
